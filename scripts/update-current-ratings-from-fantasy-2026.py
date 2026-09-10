#!/usr/bin/env python3
"""Fair multi-factor Current Super League peak-rating calibration (2026).

Balances:
  - live Fantasy rate (PPG), not season totals alone
  - research priors (xlsx OVR / previous OVR / two-season index / price)
  - injury / low games as missing data (trust prior, don't crush)
  - soft elite floors for long-term stars
  - mild team-context residual (good on bad / bad on good) — not hard ceilings

Usage:
  python scripts/update-current-ratings-from-fantasy-2026.py
  npx tsx scripts/apply-player-attrs.ts --batch data/player-attr-batch.json
"""
from __future__ import annotations

import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FANTASY_MD = ROOT / "data" / "imports" / "fantasy-super-league-players-2026.md"
XLSX = ROOT / "super_league_2026_ratings.xlsx"
OUT_BATCH = ROOT / "data" / "player-attr-batch.json"
OUT_REPORT = ROOT / "data" / "fantasy-rating-live-update-2026-report.json"

ERA_26 = {
    "Bradford Bulls",
    "Castleford Tigers",
    "Catalans Dragons",
    "Huddersfield Giants",
    "Hull FC",
    "Hull KR",
    "Leeds Rhinos",
    "Leigh Leopards",
    "St Helens",
    "Toulouse Olympique",
    "Wakefield Trinity",
    "Warrington Wolves",
    "Wigan Warriors",
    "York Knights",
}

POS = sorted(
    [
        "Full Back",
        "Winger",
        "Centre",
        "Stand Off",
        "Scrum Half",
        "Prop",
        "Hooker",
        "Second Row",
        "Loose Forward",
    ],
    key=len,
    reverse=True,
)
POS_PAT = "|".join(re.escape(p) for p in POS)
POS_GROUP = {
    "FULLBACK": "OUTSIDE",
    "WING": "OUTSIDE",
    "CENTRE": "OUTSIDE",
    "STAND_OFF": "HALVES",
    "SCRUM_HALF": "HALVES",
    "HOOKER": "HOOKER",
    "PROP": "PROP",
    "SECOND_ROW": "BACKROW",
    "LOOSE_FORWARD": "BACKROW",
}

# Season-total peer knots (secondary signal when volume is meaningful).
OUTSIDE_CURVE = [
    (0, 74),
    (300, 75),
    (500, 77),
    (700, 80),
    (850, 83),
    (1000, 86),
    (1150, 88),
    (1300, 90),
    (1450, 92),
]
HALVES_CURVE = [
    (0, 74),
    (300, 75),
    (500, 77),
    (700, 80),
    (900, 83),
    (1050, 86),
    (1200, 89),
    (1350, 91),
    (1500, 93),
]
GENERIC_CURVE = [
    (0, 74),
    (300, 75),
    (500, 77),
    (700, 80),
    (900, 83),
    (1100, 86),
    (1250, 88),
    (1400, 90),
    (1550, 92),
]
CURVE_MAP = {
    "OUTSIDE": OUTSIDE_CURVE,
    "HALVES": HALVES_CURVE,
    "HOOKER": GENERIC_CURVE,
    "PROP": GENERIC_CURVE,
    "BACKROW": GENERIC_CURVE,
    "OTHER": GENERIC_CURVE,
}

# Fantasy points-per-game → rating (primary live signal).
PPG_CURVE = [
    (12, 74),
    (20, 76),
    (26, 78),
    (32, 81),
    (36, 83),
    (40, 85),
    (44, 87),
    (48, 89),
    (52, 91),
    (58, 93),
    (68, 94),
]

# Fantasy price → implied ability prior.
PRICE_CURVE = [
    (50_000, 74),
    (70_000, 76),
    (90_000, 78),
    (105_000, 81),
    (120_000, 84),
    (135_000, 87),
    (150_000, 90),
    (165_000, 92),
    (180_000, 94),
]

# Never overwrite these intentional locks / named elites.
PROTECTED = {
    "wigan-cur-bevan-french",
    "hull-kr-cur-mikey-lewis",
    "wigan-cur-zach-eckersley",
    "warrington-cur-matty-ashton",
    "castleford-cur-daejarn-asi",
    "toulouse-cur-olly-ashall-bott",
}
MANUAL_LOCKS: dict[str, int] = {
    "bradford-cur-ryan-sutton": 84,
    "bradford-cur-loghan-lewis": 78,
    "toulouse-cur-olly-ashall-bott": 87,
}

# Known crushed-then-restored peaks for injury cases.
PRE_AUDIT_PEAK: dict[str, int] = {
    "huddersfield-cur-matty-english": 86,
    "huddersfield-cur-niall-evalds": 81,
    "wigan-cur-liam-farrell": 85,
    "hull-fc-cur-herman-eseese": 89,
    "wigan-cur-jai-field": 86,
}

# BBC Super League table — last updated 5 Sep 2026.
TABLE_POS: dict[str, int] = {
    "Wigan Warriors": 1,
    "Leeds Rhinos": 2,
    "Warrington Wolves": 3,
    "Wakefield Trinity": 4,
    "Leigh Leopards": 5,
    "Hull KR": 6,
    "St Helens": 7,
    "Toulouse Olympique": 8,
    "Catalans Dragons": 9,
    "York Knights": 10,
    "Castleford Tigers": 11,
    "Hull FC": 12,
    "Huddersfield Giants": 13,
    "Bradford Bulls": 14,
}

INJURY_RE = re.compile(
    r"injur|surgery|rehab|sidelined|unavailable|hamstring|fracture|"
    r"reduced availability|missed (most|much|games|rounds|time)|"
    r"\bankle\b|\bknee\b|\bACL\b",
    re.I,
)

MIN_GAP = 3
MAX_STEP_UP = 3
MAX_STEP_DOWN = 3


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def interp(x: float, knots: list[tuple[float, float]]) -> float:
    if x <= knots[0][0]:
        return float(knots[0][1])
    if x >= knots[-1][0]:
        return float(knots[-1][1])
    for i in range(1, len(knots)):
        x0, y0 = knots[i - 1]
        x1, y1 = knots[i]
        if x0 <= x <= x1:
            return y0 + (x - x0) / max(1e-9, x1 - x0) * (y1 - y0)
    return float(knots[-1][1])


def clamp_step(cur: int, tgt: int, up: int, down: int) -> int:
    d = tgt - cur
    if d > up:
        return cur + up
    if d < -down:
        return cur - down
    return tgt


def table_pos(club: str) -> int:
    if club in TABLE_POS:
        return TABLE_POS[club]
    for name, pos in TABLE_POS.items():
        if norm(name) == norm(club):
            return pos
    return 10


def parse_price(raw: str) -> int | None:
    m = re.search(r"([\d.]+)\s*k", raw or "", re.I)
    if not m:
        return None
    try:
        return int(float(m.group(1)) * 1000)
    except ValueError:
        return None


def parse_fantasy(text: str) -> list[dict]:
    out: list[dict] = []
    for line in text.splitlines():
        if not line.startswith("|"):
            continue
        parts = [p.strip() for p in line.strip().strip("|").split("|")]
        if len(parts) < 12:
            continue
        name_cell = parts[1]
        if name_cell in ("Name", "---") or parts[0] == "Team":
            continue
        m = re.match(rf"^(.+?)({POS_PAT})$", name_cell)
        name = m.group(1).strip() if m else name_cell
        try:
            last_rd = int(parts[-2].replace(",", ""))
            score = int(parts[-1].replace(",", ""))
        except ValueError:
            continue
        price_raw = parts[3] if len(parts) > 3 else ""
        try:
            metres = int(parts[4].replace(",", ""))
            carries = int(parts[5].replace(",", ""))
            tackles = int(parts[6].replace(",", ""))
        except (ValueError, IndexError):
            metres = carries = tackles = 0
        out.append(
            {
                "name": name,
                "norm": norm(name),
                "priceRaw": price_raw,
                "price": parse_price(price_raw),
                "metres": metres,
                "carries": carries,
                "tackles": tackles,
                "lastRd": last_rd,
                "score": score,
            }
        )
    return out


def load_xlsx() -> dict[str, dict]:
    if not XLSX.exists():
        return {}
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["All Players"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    xlsx: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        if d.get("Player"):
            xlsx[norm(str(d["Player"]))] = d
    return xlsx


def xlsx_injury(row: dict | None) -> bool:
    if not row:
        return False
    blob = " ".join(
        str(row.get(k) or "")
        for k in (
            "Rating Rationale",
            "Confidence",
            "Form Evidence",
            "Squad Status",
            "Why",
        )
    )
    return bool(INJURY_RE.search(blob))


def expected_floor(price: int | None) -> int:
    if not price:
        return 200
    if price >= 150_000:
        return 700
    if price >= 130_000:
        return 550
    if price >= 110_000:
        return 400
    if price >= 90_000:
        return 250
    return 100


def classify_availability(live: dict, xrow: dict | None) -> str:
    score = int(live.get("score") or 0)
    last_rd = int(live.get("lastRd") or 0)
    metres = int(live.get("metres") or 0)
    price = live.get("price")
    floor = expected_floor(price)
    inj = xlsx_injury(xrow)

    if inj and score < floor:
        return "injured_out"
    if last_rd == 0 and (price or 0) >= 110_000 and score < floor:
        return "injured_out"
    if metres < 40 and (price or 0) >= 110_000 and score < floor * 0.7:
        return "injured_out"
    if last_rd == 0 and (price or 0) >= 100_000:
        return "limited"
    if inj:
        return "limited"
    return "available"


def num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    return None


def estimate_games(live: dict, group: str) -> float:
    """Estimate appearances from metres / tackles / carries only."""
    metres = float(live.get("metres") or 0)
    tackles = float(live.get("tackles") or 0)
    carries = float(live.get("carries") or 0)

    norms = {
        "OUTSIDE": (95.0, 8.0, 6.0),
        "HALVES": (70.0, 16.0, 8.0),
        "HOOKER": (45.0, 38.0, 8.0),
        "PROP": (65.0, 22.0, 12.0),
        "BACKROW": (75.0, 28.0, 12.0),
        "OTHER": (70.0, 20.0, 10.0),
    }
    m_pg, t_pg, c_pg = norms.get(group, norms["OTHER"])
    estimates: list[float] = []
    if metres >= 40:
        estimates.append(metres / m_pg)
    if tackles >= 15:
        estimates.append(tackles / t_pg)
    if carries >= 12:
        estimates.append(carries / c_pg)

    if not estimates:
        return 0.0
    estimates.sort()
    return max(0.0, min(27.0, estimates[len(estimates) // 2]))


def live_weight(games: float, avail: str, score: int) -> float:
    if avail == "injured_out":
        return 0.08
    if avail == "limited":
        return 0.18 if games >= 8 else 0.10
    # High score with tiny involvement estimate → don't overweight invented PPG.
    if score >= 500 and games < 6:
        return 0.22
    if games >= 18:
        return 0.55
    if games >= 12:
        return 0.45
    if games >= 8:
        return 0.32
    if games >= 5:
        return 0.22
    return 0.12


def team_residual(ppg: float, table: int) -> float:
    """Mild ability residual: reward standouts on weak sides, trim tourists on strong sides."""
    # Expected PPG rises slightly up the table (structure / opportunity).
    if table <= 4:
        expected = 40.0
    elif table <= 7:
        expected = 36.0
    elif table <= 10:
        expected = 33.0
    else:
        expected = 30.0
    delta = ppg - expected
    # ±1.5 rating points max from team context.
    return max(-1.5, min(1.5, delta / 12.0))


def soft_elite_floor(
    pid: str,
    price: int | None,
    xrow: dict | None,
    priors: list[float],
) -> int | None:
    """Floor only for true elites / known injury restores — not every 'Star' tier."""
    tier = str((xrow or {}).get("Tier") or "").lower()
    pre = PRE_AUDIT_PEAK.get(pid)
    base = max(priors) if priors else None
    if pre is not None:
        base = max(base or 0.0, float(pre))
    if base is None:
        return None

    # Named injury restores — snap back to the known pre-crush peak.
    if pre is not None:
        return int(pre)

    if "world class" in tier or (price or 0) >= 165_000:
        if base >= 90:
            return max(90, int(round(base)) - 1)
        if base >= 88:
            return int(round(base)) - 1
        return None

    if "elite" in tier or (price or 0) >= 150_000:
        if base >= 88:
            return max(87, int(round(base)) - 2)
        return None

    # Star tier only floors when the prior itself is already high.
    if "star" in tier and base >= 87:
        return int(round(base)) - 2

    return None


def build_prior_rating(
    cur: int,
    live: dict,
    xrow: dict | None,
    pid: str,
) -> tuple[float, dict]:
    parts: list[tuple[str, float, float]] = []  # label, value, weight
    parts.append(("current", float(cur), 0.28))

    ovr = num(xrow.get("OVR")) if xrow else None
    prev = num(xrow.get("Previous OVR")) if xrow else None
    idx = num(xrow.get("Two-Season Form Index")) if xrow else None
    if ovr is not None:
        parts.append(("xlsxOvr", ovr, 0.28))
    if prev is not None:
        parts.append(("xlsxPrev", prev, 0.14))
    if idx is not None:
        parts.append(("twoSeason", idx, 0.18))

    price = live.get("price")
    if price:
        parts.append(("price", interp(price, PRICE_CURVE), 0.16))

    pre = PRE_AUDIT_PEAK.get(pid)
    if pre is not None:
        parts.append(("preAudit", float(pre), 0.20))

    tw = sum(w for _, _, w in parts) or 1.0
    blended = sum(v * w for _, v, w in parts) / tw
    detail = {lab: round(v, 1) for lab, v, _ in parts}
    detail["priorBlend"] = round(blended, 2)
    return blended, detail


def fair_target(
    *,
    pid: str,
    cur: int,
    live: dict,
    xrow: dict | None,
    club: str,
    position: str,
) -> tuple[int, dict]:
    avail = classify_availability(live, xrow)
    score = int(live.get("score") or 0)
    price = live.get("price")
    table = table_pos(club)
    group = POS_GROUP.get(position or "", "OTHER")
    games = estimate_games(live, group)
    ppg = (score / games) if games >= 4 else 0.0
    if ppg > 0:
        ppg = min(ppg, 60.0)

    # Thin samples: gently mix in proven 2025 rate when available.
    avg_2025 = num(xrow.get("2025 Avg/Game")) if xrow else None
    if ppg > 0 and avg_2025 and avg_2025 >= 18 and games < 12:
        ppg = 0.7 * ppg + 0.3 * min(avg_2025, 60.0)

    prior, prior_detail = build_prior_rating(cur, live, xrow, pid)
    lw = live_weight(games, avail, score)

    if score <= 0:
        live_r = prior
        vol_r = prior
        form = prior
    else:
        vol_r = interp(score, CURVE_MAP[group])
        if table >= 11 and games >= 12:
            vol_r -= 1.0
        elif table >= 8 and games >= 14:
            vol_r -= 0.5

        # High season total with broken/low involvement estimate → trust volume.
        unreliable_games = score >= 800 and games < 8

        if unreliable_games:
            live_r = vol_r
            form = vol_r
            lw = max(lw, 0.50)
        elif ppg > 0 and games >= 8:
            live_r = interp(ppg, PPG_CURVE)
            form = 0.55 * live_r + 0.45 * vol_r
        elif ppg > 0 and games >= 4:
            live_r = interp(ppg, PPG_CURVE)
            form = 0.35 * live_r + 0.65 * vol_r
        else:
            live_r = vol_r
            form = vol_r

    residual = team_residual(ppg, table) if ppg > 0 and games >= 8 and avail == "available" else 0.0
    form += residual

    # Injury / sparse minutes: mostly trust prior.
    if avail == "injured_out":
        target_f = 0.85 * prior + 0.15 * form
        max_down, max_up = 1, 2
    elif avail == "limited" and games < 10:
        target_f = 0.75 * prior + 0.25 * form
        max_down, max_up = 2, 2
    else:
        target_f = (1.0 - lw) * prior + lw * form
        max_down, max_up = MAX_STEP_DOWN, MAX_STEP_UP
        if target_f < prior - 4:
            max_down = min(max_down, 2)
        if target_f > prior + 5 and table >= 11:
            max_up = min(max_up, 2)

    # Don't cut strong season producers just because research priors are lower.
    if target_f < cur and score >= 850 and vol_r >= cur - 1.5:
        target_f = float(cur)

    # Low fantasy with no injury signal — don't invent upgrades.
    if score < 200 and avail == "available" and target_f > cur:
        target_f = float(cur)

    floor = soft_elite_floor(
        pid,
        price,
        xrow,
        [prior]
        + ([v for v in [num((xrow or {}).get("OVR")), num((xrow or {}).get("Previous OVR")), num((xrow or {}).get("Two-Season Form Index"))] if v is not None]),
    )
    target_f = max(72.0, min(94.0, target_f))
    if floor is not None:
        target_f = max(target_f, float(floor))
        # Restore crushed elites / injury cases in one pass (cap wild jumps).
        if cur < floor:
            span = floor - cur
            max_up = max(max_up, span if pid in PRE_AUDIT_PEAK else min(6, span))

    # Thin involvement without a rate read — keep upgrades modest unless flooring.
    if games < 4 and avail == "available" and (floor is None or cur >= (floor or 0)):
        max_up = min(max_up, 2)
        if target_f > cur + 2:
            target_f = float(cur + 2)

    # Bottom clubs: standouts can reach high-80s, not automatic 90s.
    if table >= 11:
        soft_cap = 88
        if score >= 1300 and games >= 14:
            soft_cap = 90
        # Elite floors still win (Ese'ese-class), but soft_cap binds form-only rises.
        if floor is None:
            target_f = min(target_f, float(soft_cap))
        else:
            target_f = min(max(target_f, float(floor)), max(float(soft_cap), float(floor)))

    target = int(round(target_f))
    meta = {
        "availability": avail,
        "tablePos": table,
        "gamesEst": round(games, 1),
        "ppg": round(ppg, 1),
        "liveWeight": round(lw, 2),
        "liveRate": round(live_r, 1),
        "volImplied": round(vol_r, 1),
        "form": round(form, 1),
        "teamResidual": round(residual, 2),
        "prior": round(prior, 1),
        "floor": floor,
        "maxUp": max_up,
        "maxDown": max_down,
        "fantasyScore": score,
        "price": price,
        "group": group,
        **prior_detail,
    }
    return target, meta


def match_live(name: str, live_by: dict[str, dict]) -> dict | None:
    live = live_by.get(norm(name))
    if live:
        return live
    n = norm(name)
    for k, v in live_by.items():
        if n in k or k in n:
            return v
    return None


def match_xlsx(name: str, xlsx: dict[str, dict]) -> dict | None:
    row = xlsx.get(norm(name))
    if row:
        return row
    n = norm(name)
    for k, v in xlsx.items():
        if n in k or k in n:
            return v
    return None


def main() -> None:
    if not FANTASY_MD.exists():
        raise SystemExit(f"Missing Fantasy dump: {FANTASY_MD}")

    fantasy = parse_fantasy(FANTASY_MD.read_text(encoding="utf-8", errors="replace"))
    live_by: dict[str, dict] = {}
    for row in fantasy:
        prev = live_by.get(row["norm"])
        if not prev or row["score"] > prev["score"]:
            live_by[row["norm"]] = row

    xlsx = load_xlsx()
    squads = json.loads((ROOT / "data" / "current-squads.json").read_text(encoding="utf-8"))
    if isinstance(squads, dict):
        raise SystemExit("Unexpected current-squads.json shape")
    era = [p for p in squads if (p.get("club") or p.get("team")) in ERA_26]

    ov = json.loads((ROOT / "data" / "player-rating-overrides.json").read_text(encoding="utf-8"))
    pot_ov = {k: int(v) for k, v in (ov.get("potentialOverrides") or {}).items()}

    ratings: dict[str, int] = {}
    pots: dict[str, int] = {}
    changes: list[dict] = []
    skipped: list[dict] = []
    unmatched = 0

    for p in era:
        pid = p["id"]
        name = p["name"]
        cur = int(p["peakRating"])
        club = p.get("club") or p.get("team") or ""
        live = match_live(name, live_by)
        xrow = match_xlsx(name, xlsx)

        if not live:
            unmatched += 1
            # Still allow prior-only nudge for known injury restores.
            if pid in PRE_AUDIT_PEAK and cur < PRE_AUDIT_PEAK[pid] - 1:
                to = clamp_step(cur, PRE_AUDIT_PEAK[pid], 2, 1)
                if to != cur:
                    ratings[pid] = to
                    changes.append(
                        {
                            "id": pid,
                            "name": name,
                            "club": club,
                            "from": cur,
                            "to": to,
                            "fantasyScore": 0,
                            "reason": f"no fantasy — restore pre-audit {PRE_AUDIT_PEAK[pid]}",
                        }
                    )
                    continue
            skipped.append({"id": pid, "name": name, "reason": "no fantasy row"})
            continue

        if pid in MANUAL_LOCKS:
            lock = MANUAL_LOCKS[pid]
            if cur != lock:
                ratings[pid] = lock
                changes.append(
                    {
                        "id": pid,
                        "name": name,
                        "club": club,
                        "from": cur,
                        "to": lock,
                        "fantasyScore": live["score"],
                        "reason": f"manual lock {lock}",
                    }
                )
            else:
                skipped.append({"id": pid, "name": name, "reason": "manual lock"})
            continue

        if pid in PROTECTED:
            skipped.append(
                {
                    "id": pid,
                    "name": name,
                    "reason": f"protected elite ({live['score']} fantasy)",
                    "fantasyScore": live["score"],
                }
            )
            continue

        target, meta = fair_target(
            pid=pid,
            cur=cur,
            live=live,
            xrow=xrow,
            club=club,
            position=p.get("position") or "",
        )

        if abs(cur - target) < MIN_GAP:
            skipped.append(
                {
                    "id": pid,
                    "name": name,
                    "reason": (
                        f"within +/-{MIN_GAP - 1} of fair ~{target} "
                        f"(prior {meta['prior']}, form {meta['form']}, "
                        f"games {meta['gamesEst']}, {meta['availability']})"
                    ),
                    "fantasyScore": meta["fantasyScore"],
                    "tablePos": meta["tablePos"],
                }
            )
            continue

        to = clamp_step(cur, target, meta["maxUp"], meta["maxDown"])
        if to == cur:
            continue

        ratings[pid] = to
        by = p.get("birthYear") or 0
        pot = pot_ov.get(pid)
        if by >= 2002:
            pot = max(pot or 0, to + 2)
        elif pot is not None:
            pot = max(pot, to)
        if pot is not None:
            pots[pid] = pot

        direction = "up" if to > cur else "down"
        changes.append(
            {
                "id": pid,
                "name": name,
                "club": club,
                "from": cur,
                "to": to,
                "target": target,
                "lastRd": live["lastRd"],
                **meta,
                "reason": (
                    f"fair blend prior {meta['prior']} + form {meta['form']} "
                    f"(PPG {meta['ppg']} ×{meta['gamesEst']}g, w={meta['liveWeight']}, "
                    f"{meta['availability']}, table {meta['tablePos']}, {direction})"
                ),
            }
        )

    for pid, r in list(ratings.items()):
        if pid in pots:
            pots[pid] = max(pots[pid], r)

    changes.sort(key=lambda c: (-abs(c["to"] - c["from"]), c["name"]))
    batch = {"ratings": ratings, "potentials": pots}
    report = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "method": "fair-multi-factor-2026",
        "source": str(FANTASY_MD),
        "xlsxSource": str(XLSX) if XLSX.exists() else None,
        "tableSource": "BBC Super League table (5 Sep 2026)",
        "tablePos": TABLE_POS,
        "weightsNote": (
            "Live PPG weighted by estimated games; injured/limited trust priors; "
            "soft elite floors; mild team residual ±1.5; asymmetric steps."
        ),
        "fantasyRows": len(fantasy),
        "eraPlayers": len(era),
        "xlsxPlayers": len(xlsx),
        "unmatched": unmatched,
        "ratingChanges": len(ratings),
        "potentialSets": len(pots),
        "minGap": MIN_GAP,
        "maxStepUp": MAX_STEP_UP,
        "maxStepDown": MAX_STEP_DOWN,
        "changes": changes,
        "skippedSample": skipped[:50],
        "skippedCount": len(skipped),
    }
    OUT_BATCH.write_text(json.dumps(batch, indent=2) + "\n", encoding="utf-8")
    OUT_REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")

    ups = sum(1 for c in changes if c["to"] > c["from"])
    downs = sum(1 for c in changes if c["to"] < c["from"])
    print(f"Fantasy rows: {len(fantasy)} | Era players: {len(era)} | Xlsx: {len(xlsx)}")
    print(f"Rating changes: {len(ratings)} (up {ups} / down {downs}) | Potentials: {len(pots)}")
    print(f"Skipped: {len(skipped)} | Unmatched fantasy: {unmatched}")
    print("\nName                         From   To   d  Fant  Reason")
    for c in changes[:70]:
        d = c["to"] - c["from"]
        name = c["name"].encode("ascii", "replace").decode("ascii")
        reason = c["reason"][:72].encode("ascii", "replace").decode("ascii")
        print(f"{name[:28]:28} {c['from']:4} {c['to']:4} {d:+3} {c.get('fantasyScore', 0):5}  {reason}")
    if len(changes) > 70:
        print(f"... +{len(changes) - 70} more")
    print("\nWrote", OUT_BATCH)
    print("Wrote", OUT_REPORT)


if __name__ == "__main__":
    main()
